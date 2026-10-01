import os
from datetime import datetime
from flask import Blueprint, render_template, request, redirect, url_for, flash
from flask_login import login_required
import openpyxl

from app import db
from models.product import Product
from models.supplier import Supplier
from models.customer import Customer
from models.purchase import Purchase, PurchaseItem
from models.stock import StockMovement
from models.sale import Sale, SaleItem
from models.expense import Expense
from models.partner import Partner, PartnerCapitalTransaction
from models.cash_bank import Account, CashBankTransaction

import_bp = Blueprint("import_excel", __name__)

DEFAULT_EXCEL_PATH = r"D:\Aqib Murree\Crockery_Sele_P Ragistaer.xlsx"


def safe_float(val, default=0.0):
    try:
        return float(val) if val is not None else default
    except (ValueError, TypeError):
        return default


@import_bp.route("/import-excel", methods=["GET", "POST"])
@login_required
def index():
    file_path = DEFAULT_EXCEL_PATH
    file_exists = os.path.exists(file_path)

    preview_products = []
    preview_purchases = []
    preview_sales = []
    preview_expenses = []
    preview_partner_capitals = []

    if file_exists:
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            if "Stock" in wb.sheetnames:
                for r in list(wb["Stock"].iter_rows(values_only=True)):
                    if not r or not r[0] or str(r[0]).strip() in ("Item Code", "Automatic Stock Register", "Enter each item once. Purchases and Sales will update the Current Qty automatically."):
                        continue
                    if r[0] and r[1]:
                        preview_products.append({
                            "code": str(r[0]).strip(),
                            "name": str(r[1]).strip(),
                            "unit": str(r[2] or "Pcs").strip(),
                            "opening_qty": safe_float(r[3]),
                            "cost_rate": safe_float(r[7]),
                            "reorder_level": safe_float(r[9], 10.0),
                        })

            if "Purchases" in wb.sheetnames:
                for r in list(wb["Purchases"].iter_rows(values_only=True)):
                    if not r or not r[1] or str(r[1]).strip() in ("Bill No.", "Bill No", "Purchase Register", "None"):
                        continue
                    if str(r[0] or "").strip() == "Tip:":
                        continue
                    if r[1] and r[2]:  # Bill No and Supplier
                        preview_purchases.append({
                            "date": str(r[0] or "").strip(),
                            "bill_no": str(r[1]).strip(),
                            "supplier": str(r[2]).strip(),
                            "item": str(r[3] or "").strip(),
                            "qty": safe_float(r[4]),
                            "unit": str(r[5] or "Pcs").strip(),
                            "rate": safe_float(r[6]),
                            "total": safe_float(r[7]),
                            "paid": safe_float(r[8]),
                            "balance": safe_float(r[9]),
                        })

            if "Sales" in wb.sheetnames:
                for r in list(wb["Sales"].iter_rows(values_only=True)):
                    if not r or not r[1] or str(r[1]).strip() in ("Invoice No.", "Invoice No", "Sales Register", "None"):
                        continue
                    if str(r[0] or "").strip() == "Tip:":
                        continue
                    if r[1] and r[2]:  # Invoice No and Customer
                        preview_sales.append({
                            "date": str(r[0] or "").strip(),
                            "invoice_no": str(r[1]).strip(),
                            "customer": str(r[2]).strip(),
                            "item": str(r[3] or "").strip(),
                            "qty": safe_float(r[4]),
                            "unit": str(r[5] or "Pcs").strip(),
                            "cost_rate": safe_float(r[6]),
                            "sale_rate": safe_float(r[7]),
                            "cost_total": safe_float(r[8]),
                            "total_sale": safe_float(r[9]),
                            "gross_profit": safe_float(r[10]),
                            "received": safe_float(r[11]),
                            "balance": safe_float(r[12]),
                        })

            if "Expenses" in wb.sheetnames:
                for r in list(wb["Expenses"].iter_rows(values_only=True)):
                    if not r or not r[1] or str(r[1]).strip() in ("Expense Head", "Business Expense Register", "None"):
                        continue
                    if safe_float(r[4]) > 0:
                        preview_expenses.append({
                            "date": str(r[0] or "").strip(),
                            "category": str(r[1]).strip(),
                            "description": str(r[2] or "").strip(),
                            "paid_by": str(r[3] or "Cash").strip(),
                            "amount": safe_float(r[4]),
                            "remarks": str(r[5] or "").strip(),
                        })

            if "Partners Capital" in wb.sheetnames:
                for r in list(wb["Partners Capital"].iter_rows(values_only=True)):
                    if not r or not r[1] or str(r[1]).strip() in ("Partner", "Partners Capital", "None"):
                        continue
                    if safe_float(r[3]) > 0 or safe_float(r[4]) > 0:
                        preview_partner_capitals.append({
                            "date": str(r[0] or "").strip(),
                            "partner": str(r[1]).strip(),
                            "description": str(r[2] or "").strip(),
                            "added": safe_float(r[3]),
                            "returned": safe_float(r[4]),
                            "balance": safe_float(r[5]),
                        })
        except Exception as e:
            flash(f"Error reading Excel preview: {str(e)}", "warning")

    if request.method == "POST":
        action = request.form.get("action")
        if action == "execute_import" and file_exists:
            try:
                wb = openpyxl.load_workbook(file_path, data_only=True)
                products_imported = 0
                suppliers_imported = 0
                customers_imported = 0
                purchases_imported = 0
                sales_imported = 0

                # 1. Import Products from Stock Sheet
                if "Stock" in wb.sheetnames:
                    for r in list(wb["Stock"].iter_rows(values_only=True)):
                        if not r or not r[0] or str(r[0]).strip() in ("Item Code", "Automatic Stock Register", "Enter each item once. Purchases and Sales will update the Current Qty automatically."):
                            continue
                        code = str(r[0]).strip()
                        name = str(r[1] or code).strip()
                        unit = str(r[2] or "Pcs").strip()
                        opening_qty = safe_float(r[3])
                        cost_rate = safe_float(r[7])
                        reorder = safe_float(r[9], 10.0)

                        prod = Product.query.filter_by(code=code).first()
                        if not prod:
                            prod = Product(
                                code=code,
                                name=name,
                                category="Crockery",
                                unit=unit,
                                cost_price=cost_rate,
                                sale_price=cost_rate * 1.25,
                                current_stock=opening_qty,
                                min_stock_alert=reorder,
                                is_active=True,
                            )
                            db.session.add(prod)
                            products_imported += 1
                        else:
                            prod.name = name
                            prod.unit = unit
                            prod.cost_price = cost_rate
                            prod.min_stock_alert = reorder

                db.session.commit()

                # 2. Import Purchases & Suppliers
                if "Purchases" in wb.sheetnames:
                    for r in list(wb["Purchases"].iter_rows(values_only=True)):
                        if not r or not r[1] or str(r[1]).strip() in ("Bill No.", "Bill No", "Purchase Register", "None"):
                            continue
                        if str(r[0] or "").strip() == "Tip:":
                            continue
                        date_val = str(r[0] or "").strip()
                        bill_no = str(r[1]).strip()
                        sup_name = str(r[2]).strip()
                        item_name = str(r[3] or "").strip()
                        qty = safe_float(r[4])
                        unit = str(r[5] or "Pcs").strip()
                        rate = safe_float(r[6])
                        total = safe_float(r[7])
                        paid = safe_float(r[8])
                        balance = safe_float(r[9])

                        if not sup_name or not bill_no:
                            continue

                        supplier = Supplier.query.filter_by(name=sup_name).first()
                        if not supplier:
                            supplier = Supplier(
                                name=sup_name,
                                opening_balance=0.0,
                                current_balance=0.0,
                                is_active=True,
                            )
                            db.session.add(supplier)
                            db.session.flush()
                            suppliers_imported += 1

                        if Purchase.query.filter_by(bill_no=bill_no).first():
                            continue

                        product = Product.query.filter(
                            (Product.name.ilike(f"%{item_name}%")) | (Product.code.ilike(f"%{item_name}%"))
                        ).first()
                        if not product and item_name:
                            p_code = f"ITEM-{Product.query.count() + 1:03d}"
                            product = Product(
                                code=p_code,
                                name=item_name,
                                unit=unit,
                                cost_price=rate,
                                sale_price=rate * 1.25,
                                current_stock=0.0,
                                is_active=True,
                            )
                            db.session.add(product)
                            db.session.flush()

                        p_date = datetime.utcnow()
                        if date_val:
                            for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
                                try:
                                    p_date = datetime.strptime(date_val, fmt)
                                    break
                                except ValueError:
                                    pass

                        purchase = Purchase(
                            bill_no=bill_no,
                            supplier_id=supplier.id,
                            date=p_date,
                            total_amount=total,
                            paid_amount=paid,
                            balance_amount=balance,
                            payment_status="Paid" if balance <= 0 else ("Partial" if paid > 0 else "Unpaid"),
                            notes="Imported from Excel",
                        )
                        db.session.add(purchase)
                        db.session.flush()

                        if product:
                            p_item = PurchaseItem(
                                purchase_id=purchase.id,
                                product_id=product.id,
                                quantity=qty,
                                unit=unit,
                                unit_price=rate,
                                total_price=total,
                            )
                            db.session.add(p_item)
                            product.current_stock += qty

                            mov = StockMovement(
                                product_id=product.id,
                                date=p_date,
                                movement_type="IN",
                                quantity=qty,
                                unit_cost=rate,
                                reference_type="Purchase",
                                reference_id=purchase.id,
                                notes=f"Excel Import: {bill_no}",
                            )
                            db.session.add(mov)

                        supplier.current_balance += balance
                        purchases_imported += 1

                # 3. Import Sales & Customers
                if "Sales" in wb.sheetnames:
                    for r in list(wb["Sales"].iter_rows(values_only=True)):
                        if not r or not r[1] or str(r[1]).strip() in ("Invoice No.", "Invoice No", "Sales Register", "None"):
                            continue
                        if str(r[0] or "").strip() == "Tip:":
                            continue
                        date_val = str(r[0] or "").strip()
                        inv_no = str(r[1]).strip()
                        cust_name = str(r[2]).strip()
                        item_name = str(r[3] or "").strip()
                        qty = safe_float(r[4])
                        unit = str(r[5] or "Pcs").strip()
                        cost_rate = safe_float(r[6])
                        sale_rate = safe_float(r[7])
                        qty_cost = safe_float(r[8], qty * cost_rate)
                        total_sale = safe_float(r[9], qty * sale_rate)
                        gross_profit = safe_float(r[10], total_sale - qty_cost)
                        received = safe_float(r[11])
                        balance = safe_float(r[12], total_sale - received)

                        if not cust_name or not inv_no:
                            continue

                        customer = Customer.query.filter_by(name=cust_name).first()
                        if not customer:
                            customer = Customer(
                                name=cust_name,
                                opening_balance=0.0,
                                current_balance=0.0,
                                is_active=True,
                            )
                            db.session.add(customer)
                            db.session.flush()
                            customers_imported += 1

                        if Sale.query.filter_by(invoice_no=inv_no).first():
                            continue

                        product = Product.query.filter(
                            (Product.name.ilike(f"%{item_name}%")) | (Product.code.ilike(f"%{item_name}%"))
                        ).first()
                        if not product and item_name:
                            p_code = f"ITEM-{Product.query.count() + 1:03d}"
                            product = Product(
                                code=p_code,
                                name=item_name,
                                unit=unit,
                                cost_price=cost_rate,
                                sale_price=sale_rate,
                                current_stock=0.0,
                                is_active=True,
                            )
                            db.session.add(product)
                            db.session.flush()

                        s_date = datetime.utcnow()
                        if date_val:
                            for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
                                try:
                                    s_date = datetime.strptime(date_val, fmt)
                                    break
                                except ValueError:
                                    pass

                        sale = Sale(
                            invoice_no=inv_no,
                            customer_id=customer.id,
                            date=s_date,
                            total_cost=qty_cost,
                            total_sale=total_sale,
                            gross_profit=gross_profit,
                            received_amount=received,
                            balance_amount=balance,
                            payment_status="Paid" if balance <= 0 else ("Partial" if received > 0 else "Unpaid"),
                            notes="Imported from Excel",
                        )
                        db.session.add(sale)
                        db.session.flush()

                        if product:
                            s_item = SaleItem(
                                sale_id=sale.id,
                                product_id=product.id,
                                quantity=qty,
                                unit=unit,
                                cost_rate=cost_rate,
                                sale_rate=sale_rate,
                                cost_total=qty_cost,
                                sale_total=total_sale,
                                profit=gross_profit,
                            )
                            db.session.add(s_item)
                            product.current_stock -= qty

                            mov = StockMovement(
                                product_id=product.id,
                                date=s_date,
                                movement_type="OUT",
                                quantity=qty,
                                unit_cost=cost_rate,
                                reference_type="Sale",
                                reference_id=sale.id,
                                notes=f"Excel Import: {inv_no}",
                            )
                            db.session.add(mov)

                        customer.current_balance += balance
                        sales_imported += 1

                # 4. Import Expenses
                expenses_imported = 0
                if "Expenses" in wb.sheetnames:
                    cash_acc = Account.query.filter_by(account_type="Cash").first()
                    for r in list(wb["Expenses"].iter_rows(values_only=True)):
                        if not r or not r[1] or str(r[1]).strip() in ("Expense Head", "Business Expense Register", "None"):
                            continue
                        amt = safe_float(r[4])
                        if amt <= 0:
                            continue
                        cat = str(r[1]).strip()
                        desc = str(r[2] or "").strip()
                        paid_by = str(r[3] or "Cash").strip()
                        remarks = str(r[5] or "").strip()
                        date_str = str(r[0] or "").strip()

                        try:
                            exp_date = datetime.strptime(date_str, "%d/%m/%Y")
                        except ValueError:
                            exp_date = datetime.utcnow()

                        existing_exp = Expense.query.filter_by(category=cat, amount=amt, description=desc).first()
                        if not existing_exp:
                            exp_obj = Expense(
                                date=exp_date,
                                category=cat,
                                title=desc or cat,
                                description=remarks or desc,
                                amount=amt,
                                payment_method=paid_by,
                                account_id=cash_acc.id if cash_acc else None,
                            )
                            db.session.add(exp_obj)
                            expenses_imported += 1

                # 5. Import Partners Capital
                capital_imported = 0
                if "Partners Capital" in wb.sheetnames:
                    for r in list(wb["Partners Capital"].iter_rows(values_only=True)):
                        if not r or not r[1] or str(r[1]).strip() in ("Partner", "Partners Capital", "None"):
                            continue
                        p_name = str(r[1]).strip()
                        desc = str(r[2] or "").strip()
                        added = safe_float(r[3])
                        returned = safe_float(r[4])
                        date_str = str(r[0] or "").strip()

                        if added <= 0 and returned <= 0:
                            continue

                        try:
                            p_date = datetime.strptime(date_str, "%d/%m/%Y")
                        except ValueError:
                            p_date = datetime.utcnow()

                        partner = Partner.query.filter_by(name=p_name).first()
                        if not partner:
                            # If Partner 1 is default empty partner, update name
                            p1 = Partner.query.filter_by(name="Partner 1").first()
                            if p1 and (p1.capital_balance or 0.0) == 0.0:
                                partner = p1
                                partner.name = p_name
                            else:
                                partner = Partner(
                                    name=p_name,
                                    profit_share_percentage=50.0,
                                    capital_balance=0.0,
                                    current_balance=0.0,
                                )
                                db.session.add(partner)
                            db.session.flush()

                        if added > 0:
                            existing_tx = PartnerCapitalTransaction.query.filter_by(
                                partner_id=partner.id, transaction_type="Capital Added", amount=added
                            ).first()
                            if not existing_tx:
                                cap_tx = PartnerCapitalTransaction(
                                    partner_id=partner.id,
                                    date=p_date,
                                    transaction_type="Capital Added",
                                    amount=added,
                                    description=desc,
                                    reference="Excel Import",
                                )
                                db.session.add(cap_tx)
                                partner.capital_balance = (partner.capital_balance or 0.0) + added
                                partner.current_balance = (partner.current_balance or 0.0) + added
                                capital_imported += 1

                        if returned > 0:
                            existing_ret = PartnerCapitalTransaction.query.filter_by(
                                partner_id=partner.id, transaction_type="Capital Returned", amount=returned
                            ).first()
                            if not existing_ret:
                                cap_ret = PartnerCapitalTransaction(
                                    partner_id=partner.id,
                                    date=p_date,
                                    transaction_type="Capital Returned",
                                    amount=returned,
                                    description=f"Return against: {desc}",
                                    reference="Excel Import",
                                )
                                db.session.add(cap_ret)
                                partner.capital_balance = (partner.capital_balance or 0.0) - returned
                                partner.current_balance = (partner.current_balance or 0.0) - returned
                                capital_imported += 1

                db.session.commit()
                flash(
                    f"Excel data imported successfully! Products: {products_imported}, Suppliers: {suppliers_imported}, Customers: {customers_imported}, Purchases: {purchases_imported}, Sales: {sales_imported}, Expenses: {expenses_imported}, Capital Transactions: {capital_imported}.",
                    "success",
                )
                return redirect(url_for("sales.sales_list"))

            except Exception as e:
                db.session.rollback()
                flash(f"Import failed: {str(e)}", "danger")

    return render_template(
        "import_excel.html",
        file_path=file_path,
        file_exists=file_exists,
        preview_products=preview_products,
        preview_purchases=preview_purchases,
        preview_sales=preview_sales,
        preview_expenses=preview_expenses,
        preview_partner_capitals=preview_partner_capitals,
    )
